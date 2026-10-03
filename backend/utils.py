import csv
import io
import re
import unicodedata
import uuid
from datetime import date, datetime, timezone
from decimal import Decimal
from urllib.parse import quote

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

PACKAGE_STATUSES = ["AVAILABLE", "LIMITED_SEATS", "FULL", "CLOSED", "SOLD_OUT"]
REGISTRATION_STATUSES = ["NEW", "CONTACTED", "CONSULTATION", "REGISTERED", "CANCELLED", "COMPLETED"]
DOCUMENT_STATUSES = ["NOT_STARTED", "IN_PROGRESS", "COMPLETE"]
PAYMENT_STATUSES = ["UNPAID", "DP_PAID", "PARTIALLY_PAID", "PAID", "OVERDUE"]
SEAT_COUNTING_STATUSES = {"REGISTERED", "COMPLETED"}


def now_utc():
    return datetime.now(timezone.utc)


def gen_id():
    return str(uuid.uuid4())


def slugify(text: str) -> str:
    text = unicodedata.normalize("NFKD", text or "").encode("ascii", "ignore").decode()
    text = re.sub(r"[^a-zA-Z0-9]+", "-", text).strip("-").lower()
    return text or uuid.uuid4().hex[:8]


def short_code(prefix: str) -> str:
    return f"{prefix}-{now_utc().strftime('%y%m%d')}-{uuid.uuid4().hex[:4].upper()}"


def to_json_value(v):
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    return v


def row_to_dict(obj, exclude=()):
    return {c.name: to_json_value(getattr(obj, c.name)) for c in obj.__table__.columns if c.name not in exclude}


def normalize_wa(number: str) -> str:
    digits = re.sub(r"\D", "", number or "")
    if digits.startswith("0"):
        digits = "62" + digits[1:]
    elif digits.startswith("8"):
        digits = "62" + digits
    return digits


def wa_link(number: str, message: str) -> str:
    return f"https://wa.me/{normalize_wa(number)}?text={quote(message)}"


def render_template(tpl: str, mapping: dict) -> str:
    out = tpl or ""
    for k, v in mapping.items():
        out = out.replace("{" + k + "}", str(v if v is not None else "-"))
    return out


def format_rupiah(value) -> str:
    try:
        n = int(Decimal(value or 0))
    except Exception:
        n = 0
    return "Rp" + f"{n:,}".replace(",", ".")


def format_date_id(d) -> str:
    if not d:
        return "-"
    if isinstance(d, str):
        d = date.fromisoformat(d[:10])
    months = ["Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli", "Agustus", "September", "Oktober", "November", "Desember"]
    return f"{d.day} {months[d.month - 1]} {d.year}"


def compute_package_status(pkg) -> str:
    if pkg.status_override in PACKAGE_STATUSES:
        return pkg.status_override
    if not pkg.is_active:
        return "CLOSED"
    remaining = (pkg.total_seats or 0) - (pkg.seats_filled or 0)
    if remaining <= 0:
        return "FULL"
    if remaining <= (pkg.limited_threshold or 0):
        return "LIMITED_SEATS"
    return "AVAILABLE"


def compute_payment_status(price, dp, paid) -> str:
    price = Decimal(price or 0)
    dp = Decimal(dp or 0)
    paid = Decimal(paid or 0)
    if paid <= 0:
        return "UNPAID"
    if price > 0 and paid >= price:
        return "PAID"
    if dp > 0 and paid >= dp:
        return "DP_PAID"
    return "PARTIALLY_PAID"


def rows_to_csv(rows: list[dict], columns: list[tuple[str, str]]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow([label for _, label in columns])
    for r in rows:
        writer.writerow([r.get(key, "") if r.get(key) is not None else "" for key, _ in columns])
    return ("\ufeff" + buf.getvalue()).encode("utf-8")


def rows_to_xlsx(rows: list[dict], columns: list[tuple[str, str]], sheet_name="Data") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:30]
    header_fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
    ws.append([label for _, label in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    for r in rows:
        ws.append([r.get(key, "") if r.get(key) is not None else "" for key, _ in columns])
    for i, (key, label) in enumerate(columns, start=1):
        width = max([len(str(label))] + [len(str(r.get(key, "") or "")) for r in rows[:200]])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 60)
    ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
